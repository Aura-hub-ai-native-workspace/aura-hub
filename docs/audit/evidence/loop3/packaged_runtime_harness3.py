"""Packaged-runtime proof for loop 3: bundled python + bundled backend, own token, REST only.

Ingestion fixtures are produced by the HOST's libraries in this process, so the packaged code reads files it did
not write. Search assertions look for a marker that exists only in document text, never for the query (the tool
echoes its query).
"""
import json, os, pathlib, subprocess, sys, threading, time, urllib.error, urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

X = pathlib.Path(sys.argv[1])
RES = next((X / "extract").glob("squashfs-root/usr/lib/*/resources/python"))
PY, ENTRY, BACKEND = RES / "runtime/bin/python3", RES / "scripts/serve_central_agent_api.py", RES / "backend"
HOME = X / "home"; HOME.mkdir(exist_ok=True)
PORT, TOKEN = 43791, "packaged-test-token-loop3-0123456789"
PROJ, OUTSIDE, ALPHA, BETA = (X / n for n in ("proj", "outside", "alpha", "beta"))
for d_ in (PROJ, OUTSIDE, ALPHA, BETA): d_.mkdir(exist_ok=True)

import docx, openpyxl
from pypdf import PdfWriter
from reportlab.pdfgen import canvas
from PIL import Image

d = docx.Document(); d.add_heading("Plan", 1); d.add_paragraph("host paragraph HOSTDOCX-VALUE-4417")
t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Item"; t.cell(0, 1).text = "Qty"; t.cell(1, 0).text = "valves"; t.cell(1, 1).text = "12"
d.save(str(PROJ / "host.docx"))
wb = openpyxl.Workbook(); wb.active.title = "Budget"; wb.active.append(["Item", "Cost"]); wb.active.append(["HOSTXLSX-VALUE-5528 widgets", 42])
wb.create_sheet("Notes").append(["second sheet HOSTXLSX-NOTES-9931"]); wb.save(str(PROJ / "host.xlsx"))
c = canvas.Canvas(str(PROJ / "host.pdf")); c.drawString(72, 720, "host pdf page HOSTPDF-VALUE-6603"); c.showPage(); c.drawString(72, 720, "second page text"); c.save()
c = canvas.Canvas(str(PROJ / "scan.pdf")); c.rect(72, 600, 100, 100); c.showPage(); c.save()
w = PdfWriter(clone_from=str(PROJ / "host.pdf")); w.encrypt("pw")
with open(PROJ / "locked.pdf", "wb") as fh: w.write(fh)
(PROJ / "broken.docx").write_bytes(b"PK\x03\x04 not really a docx")
(PROJ / "broken.pdf").write_bytes(b"%PDF-1.7 \xff\xfe broken")
(OUTSIDE / "secret.docx").write_bytes((PROJ / "host.docx").read_bytes())
Image.new("RGB", (8, 8), "red").save(OUTSIDE / "private.png"); Image.new("RGB", (8, 8), "blue").save(PROJ / "pic.png")
ART_OUT = OUTSIDE / "artifact-out.docx"; ART_OUT.unlink(missing_ok=True)

HITS = []
class H(BaseHTTPRequestHandler):
    def _a(self):
        HITS.append((self.command, self.path))
        if self.path.startswith("/redir"):
            self.send_response(302); self.send_header("location", self.path.split("to=", 1)[1]); self.send_header("content-length", "0"); self.end_headers(); return
        self.send_response(200); self.send_header("content-length", "2"); self.end_headers(); self.wfile.write(b"ok")
    do_GET = do_POST = do_DELETE = do_PUT = _a
    def log_message(self, *a): pass
srv = ThreadingHTTPServer(("127.0.0.1", 0), H); threading.Thread(target=srv.serve_forever, daemon=True).start()
LOCAL = f"http://127.0.0.1:{srv.server_address[1]}"

env = {"PATH": os.environ["PATH"], "PYTHONPATH": str(BACKEND), "PYTHONDONTWRITEBYTECODE": "1", "PYTHONUNBUFFERED": "1",
       "AURA_HOME": str(HOME), "AURA_INSTANCE_TOKEN": TOKEN}
log = open(X / "backend.log", "w")
child = subprocess.Popen([str(PY), str(ENTRY), str(PORT)], cwd=str(HOME), env=env, stdin=subprocess.DEVNULL, stdout=log, stderr=log)
print("child pid", child.pid)
results = []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS " if cond else "FAIL ") + name + (f" | {detail}" if detail else ""), flush=True)
def req(method, path, body=None, files=None):
    headers = {"x-aura-instance-token": TOKEN}; data = None
    if files is not None:
        boundary = "----aurabound"; parts = []
        for k, v in (body or {}).items():
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
        for k, (fn, content) in files.items():
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"; filename="{fn}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + content + b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode()); data = b"".join(parts)
        headers["content-type"] = f"multipart/form-data; boundary={boundary}"
    elif body is not None:
        data = json.dumps(body).encode(); headers["content-type"] = "application/json"
    r = urllib.request.Request(f"http://127.0.0.1:{PORT}{path}", data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(r, timeout=120) as resp: return resp.status, json.loads(resp.read() or b"null")
    except urllib.error.HTTPError as e:
        try: return e.code, json.loads(e.read() or b"null")
        except Exception: return e.code, None
def call(cap, inp, actor, project_id="proj-A"):
    ctx = {"cwd": str(PROJ), "actor": actor}
    if project_id is not None: ctx["projectId"] = project_id
    return req("POST", "/fabric/invoke", {"capabilityId": cap, "input": inp, "context": ctx})[1] or {}
det = lambda r: str(r.get("detail"))
ok_outcome = lambda r: r.get("outcome") in ("succeeded", "unverified")
try:
    for _ in range(90):
        try: urllib.request.urlopen(f"http://127.0.0.1:{PORT}/health", timeout=2); break
        except Exception: time.sleep(1)
    else: raise SystemExit("backend did not start; see backend.log")
    AGENT = {"kind": "agent", "id": "central-agent"}; HUMAN = {"kind": "human", "id": "user"}

    print("== packaged interpreter and libraries ==")
    probe = subprocess.run([str(PY), "-I", "-c", "import sys; sys.path.insert(0, r'%s'); import docx, openpyxl, reportlab, pypdf, PIL, lxml.etree; print('imports ok')" % (RES / "site-packages")],
                           capture_output=True, text=True, env={"PATH": "/usr/bin"})
    check("packaged interpreter imports all six document libraries", probe.returncode == 0 and "imports ok" in probe.stdout, (probe.stderr or probe.stdout).strip()[-90:])
    code = ("import sys\n"
            "for p in (r'%s', r'%s'): sys.path.insert(0, p)\n"
            "for m in list(sys.modules):\n"
            "    if m.split('.')[0] in ('pypdf','docx','openpyxl'): sys.modules[m] = None\n"
            "for n in ('pypdf','docx','openpyxl'): sys.modules[n] = None\n"
            "from aura.multimodal.document_engine import DocumentEngine, EngineConfig\n"
            "e = DocumentEngine(EngineConfig.from_env())\n"
            "for s in ('host.pdf','host.docx','host.xlsx'):\n"
            "    r = e.convert(r'%s/' + s); print(s, r.document_status.value, '|', r.note[:140])\n") % (BACKEND, RES / "site-packages", PROJ)
    miss = subprocess.run([str(PY), "-I", "-c", code], capture_output=True, text=True, env={"PATH": "/usr/bin"})
    lines = miss.stdout.strip().splitlines()
    check("libraries genuinely unloadable: each format reports dependency-missing naming both gaps",
          len(lines) == 3 and all("supported_dependency_missing" in l and "docling" in l.lower() and "not installed" in l.lower() for l in lines), (" || ".join(lines) or miss.stderr[-160:])[:220])

    print("== boundary ==")
    try: urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{PORT}/fabric/capabilities"), timeout=5); unauth = 200
    except urllib.error.HTTPError as e: unauth = e.code
    check("unauthenticated request refused", unauth == 401, f"status {unauth}")

    print("== F-28: DOCX, XLSX, PDF generation and ingestion ==")
    SPEC = {"title": "Report", "sections": [{"heading": "Summary", "body": "GEN-MARKER-7741 the quarter went well"},
                                            {"heading": "Table", "rows": [["Item", "Cost"], ["valves", "9"]], "list_items": ["alpha", "beta"]}]}
    for kind in ("docx", "xlsx", "pdf"):
        r = call("artifact.generate", {"spec": {**SPEC, "artifact_type": kind}, "output_path": f"gen.{kind}"}, AGENT)
        f = PROJ / f"gen.{kind}"
        check(f"generate {kind.upper()} inside the project", ok_outcome(r) and f.exists() and f.stat().st_size > 0, f"{r.get('outcome')} | {det(r)[:60]}")
    for name in ("gen.docx", "gen.xlsx", "gen.pdf", "host.docx", "host.xlsx", "host.pdf"):
        r = call("document.ingest", {"path": name}, AGENT); out = r.get("output") or {}
        check(f"ingest {name} ({'packaged-generated' if name.startswith('gen') else 'produced by the host libraries'})",
              ok_outcome(r) and out.get("chunkCount", 0) > 0 and out.get("engine") == "light", f"{r.get('outcome')} engine={out.get('engine')} chunks={out.get('chunkCount')}")
    for q, marker, label in (("quarter went well", "GEN-MARKER-7741", "generated document text"), ("HOSTDOCX", "HOSTDOCX-VALUE-4417", "host DOCX text"),
                             ("widgets", "HOSTXLSX-VALUE-5528", "host XLSX cell"), ("second sheet", "HOSTXLSX-NOTES-9931", "host XLSX second sheet"),
                             ("host pdf page", "HOSTPDF-VALUE-6603", "host PDF text")):
        check(f"search retrieves {label}", marker in json.dumps(call("knowledge.search", {"query": q}, AGENT)))
    print("== F-28 error behaviour (no false success) ==")
    for name in ("scan.pdf", "locked.pdf", "broken.docx", "broken.pdf"):
        r = call("document.ingest", {"path": name}, AGENT)
        check(f"{name} is refused", not ok_outcome(r), f"{r.get('outcome')} | {det(r)[:80]}")

    print("== F-26: knowledge is project-scoped ==")
    check("another project's tool search does not see it", "GEN-MARKER-7741" not in json.dumps(call("knowledge.search", {"query": "quarter went well"}, AGENT, project_id="proj-B")))
    none = call("knowledge.search", {"query": "quarter went well"}, AGENT, project_id=None)
    check("no project identity: refused, nothing returned", not ok_outcome(none) and "GEN-MARKER-7741" not in json.dumps(none), f"{none.get('outcome')} | {det(none)[:70]}")
    cross = call("knowledge.search", {"query": "quarter went well", "projectId": "proj-A"}, AGENT, project_id="proj-B")
    check("naming another project in the input is refused", not ok_outcome(cross) and "GEN-MARKER-7741" not in json.dumps(cross), f"{cross.get('outcome')} | {det(cross)[:70]}")
    ids = {}
    for name, folder in (("alpha", ALPHA), ("beta", BETA)):
        st, body = req("POST", "/projects", {"path": str(folder), "name": name}); ids[name] = (body or {}).get("id") or ((body or {}).get("project") or {}).get("id")
    check("two projects registered", all(ids.values()), str(ids))
    blob = (PROJ / "host.docx").read_bytes()
    check("REST ingest without a project is refused (400)", req("POST", "/documents/ingest", files={"file": ("a.docx", blob)})[0] == 400)
    check("REST ingest into an unregistered project is refused (404)", req("POST", "/documents/ingest", body={"projectId": "proj-nope"}, files={"file": ("a.docx", blob)})[0] == 404)
    st, body = req("POST", "/documents/ingest", body={"projectId": ids["alpha"]}, files={"file": ("a.docx", blob)})
    check("REST ingest of a DOCX into a registered project works", st == 200 and (body or {}).get("chunkCount", 0) > 0 and (body or {}).get("engine") == "light", f"status {st} {str(body)[:90]}")
    sa = req("GET", f"/knowledge/search?q=paragraph&projectId={ids['alpha']}")[1]; sb = req("GET", f"/knowledge/search?q=paragraph&projectId={ids['beta']}")[1]
    check("REST search: the owning project sees it, the other does not", "HOSTDOCX-VALUE-4417" in json.dumps(sa) and "HOSTDOCX-VALUE-4417" not in json.dumps(sb))
    check("REST search without a project is refused (400)", req("GET", "/knowledge/search?q=paragraph")[0] == 400)

    print("== confinement (loops 1-2 retested on this build) ==")
    check("setup: http-client node connected", (req("POST", "/environment/connect", {"id": "curl"})[1] or {}).get("connected") is True)
    r = call("sandbox.execute", {"command": "du -s .", "cwd": str(OUTSIDE)}, AGENT); check("sandbox: cwd outside refused", not ok_outcome(r) and "project" in det(r).lower(), det(r)[:70])
    r = call("sandbox.execute", {"command": "du -s .", "cwd": "."}, AGENT); check("sandbox: same command inside works", ok_outcome(r), det(r)[:50])
    r = call("sandbox.execute", {"command": "id"}, AGENT); check("sandbox: off-allow-list program refused", not ok_outcome(r) and "allow-list" in det(r))
    r = call("document.ingest", {"path": str(OUTSIDE / "secret.docx")}, AGENT); check("ingest: file outside the project refused", not ok_outcome(r) and "project" in det(r).lower(), det(r)[:70])
    r = call("artifact.generate", {"spec": {**SPEC, "artifact_type": "docx"}, "output_path": str(ART_OUT)}, AGENT)
    check("artifact: output outside the project refused, nothing written", not ok_outcome(r) and not ART_OUT.exists(), det(r)[:70])
    r = call("artifact.generate", {"spec": {**SPEC, "artifact_type": "docx"}, "output_path": "inside.docx"}, AGENT)
    check("artifact: output INSIDE the project works (previously blocked)", ok_outcome(r) and (PROJ / "inside.docx").exists(), f"{r.get('outcome')} | {det(r)[:60]}")
    for kind in ("docx", "pdf"):
        r = call("artifact.generate", {"spec": {"artifact_type": kind, "title": "T", "sections": [{"body": "x", "image_path": str(OUTSIDE / "private.png")}]}, "output_path": f"img-out.{kind}"}, AGENT)
        check(f"artifact {kind}: an image outside the project is refused, nothing written", not ok_outcome(r) and "image" in det(r).lower() and not (PROJ / f"img-out.{kind}").exists(), det(r)[:70])
        r = call("artifact.generate", {"spec": {"artifact_type": kind, "title": "T", "sections": [{"body": "x", "image_path": "pic.png"}]}, "output_path": f"img-in.{kind}"}, AGENT)
        check(f"artifact {kind}: an image inside the project is embedded", ok_outcome(r) and (PROJ / f"img-in.{kind}").exists(), det(r)[:60])
    HITS.clear(); check("http: agent GET to a local service works", ok_outcome(call("http.request", {"url": f"{LOCAL}/x", "method": "GET"}, AGENT)) and HITS == [("GET", "/x")])
    HITS.clear(); bad = [call("http.request", {"url": f"{LOCAL}/x", "method": m, "body": "{}"}, AGENT) for m in ("POST", "PUT", "PATCH", "DELETE")]
    check("http: agent POST/PUT/PATCH/DELETE to local refused and never sent", all(not ok_outcome(b) and "GET or HEAD" in det(b) for b in bad) and HITS == [])
    HITS.clear(); check("http: a user's PUT is allowed", ok_outcome(call("http.request", {"url": f"{LOCAL}/x", "method": "PUT", "body": "{}"}, HUMAN)) and HITS == [("PUT", "/x")])
    for who, actor in (("agent", AGENT), ("user", HUMAN)):
        r = call("http.request", {"url": "http://169.254.169.254/latest/meta-data/", "method": "GET"}, actor); check(f"http: metadata refused for the {who}", not ok_outcome(r) and "metadata" in det(r))
    r = call("http.request", {"url": f"{LOCAL}/redir?to=http://169.254.169.254/", "method": "GET"}, HUMAN); check("http: redirect to metadata refused", not ok_outcome(r) and "metadata" in det(r))

    print("== F-24 backend: Stop on an already-finished automation run ==")
    wf = req("POST", "/workflows", {"name": "W", "nodes": [{"id": "d", "type": "delay", "x": 0, "y": 0, "config": {"ms": 1}}], "edges": []})[1]
    rule = req("POST", "/automation/rules", {"name": "M", "enabled": True, "trigger": {"type": "manual"}, "conditions": [], "chain": [{"id": "a1", "action": "run-workflow", "label": "go", "config": {"workflowId": wf["id"]}, "continueOnError": False}]})[1]
    run = req("POST", f"/automation/rules/{rule['id']}/run", {"projectId": "default"})[1]
    for _ in range(50):
        if (req("GET", f"/automation/rules/{rule['id']}/runs/{run['id']}")[1] or {}).get("status") == "completed": break
        time.sleep(0.2)
    st, body = req("POST", f"/automation/rules/{rule['id']}/runs/{run['id']}/cancel")
    check("finished run: cancel answers 409 'already completed'", st == 409 and "already completed" in str(body), f"{st} {body}")
    check("unknown run: cancel answers 404", req("POST", f"/automation/rules/{rule['id']}/runs/run-nope/cancel")[0] == 404)

    print("== audit ==")
    trail = HOME / "audit/trail.jsonl"
    rows = [json.loads(l) for l in trail.read_text().splitlines()] if trail.exists() else []
    kinds = {(x.get("capabilityId"), (x.get("actor") or {}).get("kind"), x.get("outcome")) for x in rows if x.get("capabilityId") == "http.request"}
    check("audit records the actor kind for http.request", any(k[1] == "human" and k[2] == "succeeded" for k in kinds) and any(k[1] == "agent" and k[2] == "succeeded" for k in kinds))
finally:
    child.terminate()
    try: child.wait(timeout=10)
    except Exception: child.kill()
    srv.shutdown()

print("== F-23: the PACKAGED governance hook, run by the packaged interpreter ==")
hook = BACKEND / "aura/governance/claude_hook.py"
LOGF = X / "hook-actions.jsonl"
def run_hook(stdin, scope):
    LOGF.unlink(missing_ok=True)
    args = [str(PY), str(hook)] + ([str(scope)] if scope is not None else [])
    p = subprocess.run(args, input=stdin, capture_output=True, text=True, env={"PATH": "/usr/bin", "AURA_ACTION_LOG": str(LOGF), "PYTHONPATH": str(BACKEND)}, timeout=60)
    ev = [json.loads(l)["decision"] for l in LOGF.read_text().splitlines()] if LOGF.exists() else []
    return p.returncode, ev
SC = X / "scope.json"; good = {"taskId": "t", "nodeId": "n", "invocationId": "i", "attempt": "1", "cwd": "/repo", "scopePaths": ["src"], "auraRoot": str(BACKEND)}
IN = json.dumps({"tool_name": "Write", "tool_input": {"file_path": "src/a.py"}}); OUT = json.dumps({"tool_name": "Write", "tool_input": {"file_path": "billing/p.py"}})
SC.write_text(json.dumps(good))
check("hook: valid scope allows inside", run_hook(IN, SC) == (0, ["ALLOW"]))
check("hook: valid scope denies outside", run_hook(OUT, SC) == (2, ["DENY"]))
check("hook: genuinely missing scope file is allowed and recorded UNSCOPED", run_hook(IN, X / "never.json") == (0, ["UNSCOPED"]))
for label, content in (("truncated", "{trunc"), ("empty", ""), ("array", "[]")):
    SC.write_text(content); check(f"hook: {label} scope file denies, recorded DENY", run_hook(IN, SC) == (2, ["DENY"]))
SC.write_text(json.dumps({**good, "scopePaths": "src"})); check("hook: invalid field type denies", run_hook(IN, SC) == (2, ["DENY"]))
SC.write_text(json.dumps(good)); check("hook: garbage stdin denies", run_hook("{nope", SC)[0] == 2)
sd = X / "scope-dir"; sd.mkdir(exist_ok=True); check("hook: a directory in place of the scope file denies", run_hook(IN, sd) == (2, ["DENY"]))
cmd = f"python3 {hook} {SC}; rc=$?; if [ $rc -ne 0 ] && [ $rc -ne 2 ]; then exit 2; fi; exit $rc"
p = subprocess.run(["/bin/sh", "-c", cmd], input=IN, capture_output=True, text=True, env={"PATH": "/nonexistent"})
check("hook launch: a missing interpreter becomes a deny (exit 2, not 127)", p.returncode == 2, f"exit {p.returncode}")

print(f"RESULT {sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
