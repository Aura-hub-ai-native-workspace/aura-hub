"""Real-Docling integration tests (gated, never silently passed).

Gate: the Docling package must be installed AND the model cache
provisioned. Otherwise every test in this module SKIPS with an explicit
reason — a skip is reported as a skip, never as a pass.

Run inside the isolated docling venv after provisioning:
    /mnt/storage/aura-docling-venv/bin/python -m pytest \\
        tests/integration/test_docling_engine.py -q
    python scripts/provision-docling-models.py   # first-time model cache
"""

import os
import socket
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from aura.multimodal.docling_engine import (  # noqa: E402
    DoclingEngine,
    detect_ocr_engines,
    models_cached,
)
from aura.multimodal.doctypes import (  # noqa: E402
    DocumentConversionRequest,
    DocumentStatus,
    IsolationMode,
    OcrEngine,
)
from aura.multimodal.document_engine import (  # noqa: E402
    DocumentEngine,
    EngineConfig,
)

DOCLING_HERE = DoclingEngine().is_installed()
CACHE_READY = models_cached()
OCR_READY = bool(detect_ocr_engines())

needs_docling = pytest.mark.skipif(
    not (DOCLING_HERE and CACHE_READY),
    reason=f"needs docling+models (installed={DOCLING_HERE}, "
           f"cached={CACHE_READY}) — run provision-docling-models.py")
needs_ocr = pytest.mark.skipif(
    not (DOCLING_HERE and CACHE_READY and OCR_READY),
    reason=f"needs local OCR (installed={DOCLING_HERE}, "
           f"cached={CACHE_READY}, ocr={detect_ocr_engines()})")

KNOWN_LINE = "AURA OCR PROBE LINE 42"


# ----------------------------------------------------------------------
# fixture builders (all synthetic, all local)
# ----------------------------------------------------------------------

def _pdf_normal(path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story = [Paragraph("Provisioned Test Report", styles["Heading1"]),
             Paragraph("First section body with searchable token "
                       "ZEBRAHORN.", styles["Normal"]),
             Spacer(1, 12),
             Paragraph("Data Table", styles["Heading2"]),
             Table([["name", "value"], ["alpha", "1"], ["beta", "2"]],
                   style=TableStyle([
                       ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                       ("GRID", (0, 0), (-1, -1), 1, colors.black)])),
             Paragraph("Second section closes the document.",
                       styles["Normal"])]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def _text_image(path, text, fmt="PNG", size=(800, 200)):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((20, 80), text, fill="black")
    img.save(str(path), fmt)


def _pdf_scanned(path, pages=2):
    from PIL import Image

    frames = []
    for _ in range(pages):
        tmp = path.parent / "_scan_tmp.png"
        _text_image(tmp, KNOWN_LINE, size=(1200, 400))
        frames.append(Image.open(tmp))
    first, rest = frames[0], frames[1:]
    first.save(str(path), "PDF", save_all=True, append_images=rest)
    for frame in frames:
        frame.close()


def _docx(path):
    import docx

    doc = docx.Document()
    doc.core_properties.title = "Synthetic Docx"
    doc.core_properties.author = "AURA Tests"
    doc.add_heading("Docx Heading", level=1)
    doc.add_paragraph("Docx body paragraph with token QUAILBRUSH.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "h1", "h2"
    table.cell(1, 0).text, table.cell(1, 1).text = "c1", "c2"
    doc.save(str(path))


def _xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SheetOne"
    ws.append(["id", "label"])
    ws.append([1, "first"])
    ws2 = wb.create_sheet("SheetTwo")
    ws2.append(["code", "desc"])
    ws2.append(["X9", "second sheet token WOMBATRY"])
    wb.save(str(path))


def _pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide Title Token NIGHTJAR"
    slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8),
                             Inches(1)).text_frame.text = "Bullet body text."
    rows, cols = 2, 2
    graphic = slide.shapes.add_table(rows, cols, Inches(1), Inches(3),
                                     Inches(4), Inches(1))
    graphic.table.cell(0, 0).text = "t1"
    graphic.table.cell(1, 1).text = "t2"
    prs.save(str(path))


def _odt(path):
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, P

    doc = OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text="Odt Heading"))
    doc.text.addElement(P(text="Odt body with token VELVETANT."))
    table = Table()
    table.addElement(TableColumn())
    table.addElement(TableColumn())
    row = TableRow()
    for value in ("k1", "k2"):
        cell = TableCell()
        cell.addElement(P(text=value))
        row.addElement(cell)
    table.addElement(row)
    doc.text.addElement(table)
    doc.save(str(path))


def _ods(path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name="SheetOne")
    for values in (("id", "label"), ("7", "ods token MARMOTINK")):
        row = TableRow()
        for value in values:
            cell = TableCell()
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    doc.spreadsheet.addElement(table)
    doc.save(str(path))


def _soffice_convert(src, fmt, outdir):
    """Convert via LibreOffice; returns output path or None when absent."""
    import shutil
    import subprocess

    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if soffice is None:
        return None
    proc = subprocess.run(
        [soffice, "--headless", "--convert-to", fmt,
         "--outdir", str(outdir), str(src)],
        capture_output=True, text=True, timeout=180)
    candidate = outdir / (src.stem + f".{fmt}")
    if proc.returncode == 0 and candidate.is_file():
        return candidate
    return None


def _epub(path):
    container = ('<?xml version="1.0"?><container version="1.0" '
                 'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="content.opf" '
                 'media-type="application/oebps-package+xml"/>'
                 "</rootfiles></container>")
    opf = ('<?xml version="1.0"?><package version="3.0" '
           'xmlns="http://www.idpf.org/2007/opf" unique-identifier="id">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           "<dc:title>Synthetic Epub</dc:title></metadata>"
           '<manifest><item id="c" href="c.xhtml" '
           'media-type="application/xhtml+xml"/></manifest>'
           '<spine><itemref idref="c"/></spine></package>')
    xhtml = ("<html xmlns='http://www.w3.org/1999/xhtml'><head><title>T"
             "</title></head><body><h1>Chapter One</h1><p>Epub body with "
             "token LARKSPUR.</p></body></html>")
    with zipfile.ZipFile(str(path), "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("content.opf", opf)
        archive.writestr("c.xhtml", xhtml)


def _eml(path):
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["Subject"] = "Synthetic Eml Subject"
    msg["From"] = "tests@example.local"
    msg["To"] = "aura@example.local"
    msg.set_content("Eml body carrying token FOXGLOVE.")
    path.write_bytes(msg.as_bytes())


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory):
    root = tmp_path_factory.mktemp("docling_fixtures")
    out = {}
    _pdf_normal(root / "normal.pdf")
    out["pdf"] = str(root / "normal.pdf")
    _pdf_scanned(root / "scanned.pdf")
    out["scanned"] = str(root / "scanned.pdf")
    _docx(root / "doc.docx")
    out["docx"] = str(root / "doc.docx")
    _xlsx(root / "book.xlsx")
    out["xlsx"] = str(root / "book.xlsx")
    _pptx(root / "deck.pptx")
    out["pptx"] = str(root / "deck.pptx")
    _text_image(root / "line.png", KNOWN_LINE)
    out["png"] = str(root / "line.png")
    _text_image(root / "line.jpg", KNOWN_LINE, fmt="JPEG")
    out["jpg"] = str(root / "line.jpg")
    _text_image(root / "line.tiff", KNOWN_LINE, fmt="TIFF")
    out["tiff"] = str(root / "line.tiff")
    (root / "page.html").write_text(
        "<html><head><title>HT</title></head><body><h1>Html Heading</h1>"
        "<p>Html body token GECKOCLAW.</p></body></html>", encoding="utf-8")
    out["html"] = str(root / "page.html")
    (root / "notes.md").write_text("# Md Heading\n\nMd body INDIGOFINCH.\n",
                                   encoding="utf-8")
    out["md"] = str(root / "notes.md")
    (root / "rows.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    out["csv"] = str(root / "rows.csv")
    _odt(root / "doc.odt")
    out["odt"] = str(root / "doc.odt")
    _ods(root / "sheet.ods")
    out["ods"] = str(root / "sheet.ods")
    # Legacy + ODP fixtures via a real LibreOffice conversion (honest
    # end-to-end: genuine producer files, not hand-mocked bytes).
    from pathlib import Path as _Path
    legacy_doc = _soffice_convert(_Path(root / "doc.docx"), "doc", root)
    out["doc"] = str(legacy_doc) if legacy_doc else ""
    legacy_odp = _soffice_convert(_Path(root / "deck.pptx"), "odp", root)
    out["odp"] = str(legacy_odp) if legacy_odp else ""
    _epub(root / "book.epub")
    out["epub"] = str(root / "book.epub")
    _eml(root / "mail.eml")
    out["eml"] = str(root / "mail.eml")
    (root / "corrupt.pdf").write_bytes(os.urandom(2048))
    out["corrupt"] = str(root / "corrupt.pdf")
    (root / "archive.zip").write_bytes(b"PK\x03\x04not-a-real-zip")
    out["zip"] = str(root / "archive.zip")
    return out


@pytest.fixture(scope="module")
def engine():
    return DocumentEngine(EngineConfig(timeout_s=600.0))


# ----------------------------------------------------------------------
# format coverage
# ----------------------------------------------------------------------

@needs_docling
def test_pdf_normal_structure(engine, fixtures):
    result = engine.convert(fixtures["pdf"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "ZEBRAHORN" in result.text
    assert any("Provisioned Test Report" in h for h in result.headings)
    assert result.page_count >= 1
    assert result.tables, "expected at least one structured table"
    flat = " ".join(" ".join(r) for t in result.tables for r in t.rows)
    assert "alpha" in flat and "beta" in flat


@needs_ocr
def test_pdf_scanned_ocr(engine, fixtures):
    result = engine.convert(fixtures["scanned"])
    assert result.document_status in (DocumentStatus.SUCCESS,
                                      DocumentStatus.PARTIAL), result.note
    assert result.ocr_used, "expected OCR metadata on a scanned PDF"
    assert "AURA" in result.text and "42" in result.text


@needs_docling
def test_docx_paragraphs_tables_metadata(engine, fixtures):
    result = engine.convert(fixtures["docx"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "QUAILBRUSH" in result.text
    assert result.metadata.get("title") == "Synthetic Docx" or \
        "Docx Heading" in result.text


@needs_docling
def test_xlsx_multiple_sheets(engine, fixtures):
    result = engine.convert(fixtures["xlsx"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "first" in result.text
    assert "WOMBATRY" in result.text


@needs_docling
def test_pptx_slides_tables(engine, fixtures):
    result = engine.convert(fixtures["pptx"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "NIGHTJAR" in result.text


@needs_ocr
@pytest.mark.parametrize("key", ["png", "jpg", "tiff"])
def test_image_ocr(engine, fixtures, key):
    result = engine.convert(fixtures[key])
    assert result.document_status in (DocumentStatus.SUCCESS,
                                      DocumentStatus.PARTIAL), result.note
    assert "AURA" in result.text and "42" in result.text


@needs_docling
def test_html(engine, fixtures):
    result = engine.convert(fixtures["html"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "GECKOCLAW" in result.text


@needs_docling
def test_markdown_csv(engine, fixtures):
    md = engine.convert(fixtures["md"])
    assert md.document_status == DocumentStatus.SUCCESS, md.note
    assert "INDIGOFINCH" in md.text
    csv = engine.convert(fixtures["csv"])
    assert csv.document_status == DocumentStatus.SUCCESS, csv.note
    assert "1,2" in csv.text.replace(" ", "").replace("|", ",")


@needs_docling
def test_odt(engine, fixtures):
    result = engine.convert(fixtures["odt"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "VELVETANT" in result.text


@needs_docling
def test_ods(engine, fixtures):
    result = engine.convert(fixtures["ods"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "MARMOTINK" in result.text


@needs_docling
def test_legacy_doc_via_libreoffice(engine, fixtures):
    if not fixtures["doc"]:
        pytest.skip("LibreOffice conversion unavailable; "
                    "legacy .doc path untested, not passed")
    result = engine.convert(fixtures["doc"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "QUAILBRUSH" in result.text


@needs_docling
def test_odp_via_libreoffice(engine, fixtures):
    if not fixtures["odp"]:
        pytest.skip("LibreOffice conversion unavailable; "
                    "ODP path untested, not passed")
    result = engine.convert(fixtures["odp"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "NIGHTJAR" in result.text


@needs_docling
def test_epub(engine, fixtures):
    result = engine.convert(fixtures["epub"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "LARKSPUR" in result.text


@needs_docling
def test_eml(engine, fixtures):
    result = engine.convert(fixtures["eml"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "FOXGLOVE" in result.text


@needs_docling
def test_msg_capability_honest(engine):
    caps = engine.capabilities()
    # .msg support varies by installed backend; either routing or an
    # honest DEPENDENCY_MISSING is acceptable — claiming is not.
    if caps.supports(".msg"):
        assert caps.formats[".msg"].startswith("docling:")
    else:
        assert ".msg" not in caps.formats


# ----------------------------------------------------------------------
# failure behavior
# ----------------------------------------------------------------------

@needs_docling
def test_archive_rejected(engine, fixtures):
    result = engine.convert(fixtures["zip"])
    assert result.document_status == DocumentStatus.UNSUPPORTED
    assert "archive" in result.note


@needs_docling
def test_corrupt_pdf(engine, fixtures):
    result = engine.convert(fixtures["corrupt"])
    assert result.document_status == DocumentStatus.CORRUPT
    assert result.text == ""


@needs_docling
def test_encrypted_pdf_refused(tmp_path, engine):
    pytest.importorskip("pypdf")
    from pypdf import PdfReader, PdfWriter

    base = tmp_path / "base.pdf"
    _pdf_normal(base)
    locked = tmp_path / "locked.pdf"
    writer = PdfWriter()
    for page in PdfReader(str(base)).pages:
        writer.add_page(page)
    writer.encrypt("secret")
    with open(locked, "wb") as handle:
        writer.write(handle)
    result = engine.convert(str(locked))
    assert result.document_status == DocumentStatus.CORRUPT
    assert "encrypt" in result.note.lower() or "password" in \
        result.note.lower() or result.errors


@needs_docling
def test_oversized_rejected(fixtures):
    tiny = DocumentEngine(EngineConfig(max_file_bytes=10))
    result = tiny.convert(fixtures["pdf"])
    assert result.status == "error"
    assert "exceeds" in result.note


@needs_docling
def test_page_limit_enforced(fixtures):
    engine = DocumentEngine(EngineConfig(timeout_s=600.0))
    req = DocumentConversionRequest(path=fixtures["pdf"], max_pages=1)
    result = engine.convert(fixtures["pdf"], req)
    assert result.document_status in (DocumentStatus.SUCCESS,
                                      DocumentStatus.PARTIAL)
    assert result.page_count <= 1


@needs_docling
def test_symlink_escape_rejected(tmp_path, fixtures):
    root = tmp_path / "root"
    root.mkdir()
    link = root / "evil.pdf"
    link.symlink_to(fixtures["pdf"])
    engine = DocumentEngine(EngineConfig(confinement_root=str(tmp_path)))
    outside_link = tmp_path / "outside.pdf"
    outside_link.symlink_to(fixtures["pdf"])
    result = engine.convert(str(outside_link))
    assert result.document_status == DocumentStatus.CORRUPT


@needs_docling
def test_missing_ocr_engine_honest(tmp_path, fixtures, monkeypatch):
    import aura.multimodal.docling_engine as backend

    monkeypatch.setattr(backend, "detect_ocr_engines", lambda: [])
    engine = DocumentEngine(EngineConfig(
        ocr_engine=OcrEngine.RAPIDOCR, timeout_s=600.0))
    req = DocumentConversionRequest(path=fixtures["scanned"],
                                    ocr_engine=OcrEngine.RAPIDOCR)
    result = engine.docling.convert(req)
    assert result.document_status == DocumentStatus.OCR_FAILED
    assert "rapidocr" in result.note.lower()


@needs_docling
def test_process_isolation_success(tmp_path, engine, fixtures):
    req = DocumentConversionRequest(path=fixtures["md"],
                                    isolation=IsolationMode.PROCESS,
                                    timeout_s=600.0)
    result = engine.convert(fixtures["md"], req)
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert "INDIGOFINCH" in result.text


@needs_docling
def test_offline_conversion_no_downloads(engine, fixtures, monkeypatch):
    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("AURA_DOCS_OFFLINE", "1")
    result = engine.convert(fixtures["md"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note


@needs_docling
def test_zero_egress_conversion(engine, fixtures, monkeypatch):
    real_socket = socket.socket

    class _BlockedSocket:
        def __init__(self, *args, **kwargs):
            raise AssertionError("network egress during conversion")

    monkeypatch.setattr(socket, "socket", _BlockedSocket)
    try:
        result = engine.convert(fixtures["md"])
    finally:
        monkeypatch.setattr(socket, "socket", real_socket)
    assert result.document_status == DocumentStatus.SUCCESS, result.note


@needs_docling
def test_concurrent_batch_bounded(engine, fixtures):
    batch_engine = DocumentEngine(EngineConfig(max_workers=1,
                                               timeout_s=600.0))
    paths = [fixtures["md"], fixtures["csv"], fixtures["html"]]
    results = batch_engine.convert_batch(paths)
    assert len(results) == 3
    assert all(r.document_status == DocumentStatus.SUCCESS for r in results)


@needs_docling
def test_chunk_correctness(engine, fixtures):
    result = engine.convert(fixtures["pdf"])
    assert result.document_status == DocumentStatus.SUCCESS, result.note
    assert result.chunks, "expected chunks for KB ingestion"
    assert all(c.end_char > c.start_char for c in result.chunks)
    assert all(c.text.strip() for c in result.chunks)
    # Tables stay structured: cell text must not depend on chunk flattening.
    assert result.tables and result.tables[0].num_cols >= 2
    row_text = " ".join(" ".join(r) for r in result.tables[0].rows)
    assert "name" in row_text and "value" in row_text


@needs_docling
def test_kb_extraction_compat(engine, fixtures):
    result = engine.convert(fixtures["pdf"])
    extraction = result.to_extraction_result()
    assert extraction.status.value == result.status
    assert extraction.text == result.text
    assert extraction.page_count == result.page_count
