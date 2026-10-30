"""LightOfficeEngine — PDF, DOCX and XLSX text extraction without Docling.

Docling is a multi-gigabyte optional stack (torch and a model cache) and is
not shipped with the application. Before this tier existed, every DOCX, XLSX
and PDF was answered "docling package not installed", so the document features
did not work in an installed AURA at all. This engine reads those three formats
with small pure-Python libraries (`pypdf`, `python-docx`, `openpyxl`, the
`[multimodal]` and `[artifacts]` extras, which the packaged application now
ships).

It is a FALLBACK. DocumentEngine tries Docling first and uses this only when
Docling reports itself missing, so a machine that has Docling keeps its
layout-aware conversion exactly as before.

What it does not do, and says so instead of guessing: no OCR (a PDF with no
text layer is OCR_REQUIRED), no layout or table reconstruction beyond rows and
cells, no formulas (cached values only), no password-protected files.

Safety, because these are parsers run on files an agent may point at: a source
size cap, a cap on the unpacked size of DOCX/XLSX containers (a zip bomb is
refused before a parser sees it), and per-format budgets. Nothing is imported
unless needed; a missing library is a typed result naming the extra to install,
never a crash and never an empty "success".
"""

from __future__ import annotations

import mimetypes
import time
import zipfile
from pathlib import Path

from .doctypes import (
    ConversionError,
    DocumentConversionRequest,
    DocumentConversionResult,
    DocumentStatus,
)
from .native import NativeTextEngine
from .secure import (
    PathSecurityError,
    file_identity,
    file_size,
    resolve_secure,
    verify_same_file,
)

LIGHT_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx"})

#: Refuse sources larger than this, and containers that unpack to more.
MAX_SOURCE_BYTES = 100 * 1024 * 1024
MAX_UNPACKED_BYTES = 300 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_XLSX_ROWS_PER_SHEET = 20_000

_MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

_EXTRA = {
    ".pdf": ("pypdf", "pypdf", "aura-backend[multimodal]"),
    ".docx": ("docx", "python-docx", "aura-backend[multimodal]"),
    ".xlsx": ("openpyxl", "openpyxl", "aura-backend[artifacts]"),
}


class LightOfficeEngine:
    """Small-library fallback extraction for .pdf, .docx and .xlsx."""

    name = "light"

    def __init__(self, max_chars: int = 500_000) -> None:
        self._max_chars = max_chars
        self._chunker = NativeTextEngine(max_chars=max_chars)

    def claims(self, suffix: str) -> bool:
        return suffix.lower() in LIGHT_SUFFIXES

    @staticmethod
    def library_available(suffix: str) -> bool:
        """Whether this format's library can really be imported right now."""
        module = _EXTRA[suffix.lower()][0]
        try:
            __import__(module)
        except ImportError:
            return False
        return True

    # -- conversion ----------------------------------------------------
    def convert(self, request: DocumentConversionRequest) -> DocumentConversionResult:
        started = time.monotonic()
        result = DocumentConversionResult(path=request.path, engine=self.name)
        try:
            resolved = resolve_secure(request.path)
        except PathSecurityError as exc:
            return self._fail(result, DocumentStatus.CORRUPT, f"path rejected: {exc}", "policy", started)

        suffix = resolved.suffix.lower()
        if not self.claims(suffix):
            result.status = "skipped"
            result.document_status = DocumentStatus.UNSUPPORTED
            result.note = f"unsupported file type: {suffix or '(no extension)'}"
            return result
        result.mime_type = _MIME.get(suffix) or (mimetypes.guess_type(str(resolved))[0] or "")

        module, library, extra = _EXTRA[suffix]
        if not self.library_available(suffix):
            result.status = "skipped"
            result.document_status = DocumentStatus.SUPPORTED_DEPENDENCY_MISSING
            result.note = f"{library} not installed (pip install {extra})"
            result.timings_ms = {"total": (time.monotonic() - started) * 1000.0}
            return result

        try:
            size = file_size(resolved)
            identity = file_identity(resolved)
        except OSError as exc:
            return self._fail(result, DocumentStatus.CORRUPT, f"read failed: {exc}"[:200], "source_unavailable", started)
        if size > MAX_SOURCE_BYTES:
            return self._fail(result, DocumentStatus.CORRUPT,
                              f"file is larger than the {MAX_SOURCE_BYTES // (1024 * 1024)} MB limit", "policy", started)
        if suffix in (".docx", ".xlsx"):
            refused = self._container_refusal(resolved)
            if refused:
                return self._fail(result, DocumentStatus.CORRUPT, refused, "policy", started)

        limit = request.max_chars or self._max_chars
        try:
            if suffix == ".pdf":
                text, pages, note, status = self._pdf(resolved, limit)
            elif suffix == ".docx":
                text, pages, note, status = self._docx(resolved, limit)
            else:
                text, pages, note, status = self._xlsx(resolved, limit)
        except _Refused as exc:
            return self._fail(result, exc.status, str(exc), exc.category, started)
        except Exception as exc:  # noqa: BLE001 — a corrupt file is a typed result
            return self._fail(result, DocumentStatus.CORRUPT,
                              f"{library} could not read this file: {type(exc).__name__}: {exc}"[:240],
                              "corrupt", started)

        if not verify_same_file(resolved, identity):
            return self._fail(result, DocumentStatus.CORRUPT,
                              "file changed during read; refusing half-read content", "internal", started)

        result.metadata = {"sourceBytes": size, "extractor": library}
        result.page_count = pages
        result.text = text
        result.char_count = len(text)
        result.note = note
        if status is not None:
            result.status, result.document_status = status
        elif text.strip():
            result.status, result.document_status = "ok", DocumentStatus.SUCCESS
        else:
            result.status, result.document_status = "empty", DocumentStatus.PARTIAL
            result.note = result.note or "file readable but contains no text"
        if request.want_chunks and text.strip():
            result.chunks = self._chunker.chunk_text(text)
        result.timings_ms = {"total": (time.monotonic() - started) * 1000.0}
        return result

    # -- guards --------------------------------------------------------
    @staticmethod
    def _container_refusal(path: Path) -> str | None:
        """Refuse a DOCX/XLSX that is not a zip or that unpacks absurdly."""
        try:
            with zipfile.ZipFile(path) as archive:
                unpacked = sum(info.file_size for info in archive.infolist())
        except (zipfile.BadZipFile, OSError):
            return "not a valid Office file (the container is unreadable)"
        if unpacked > MAX_UNPACKED_BYTES:
            return (f"refused: the file unpacks to {unpacked // (1024 * 1024)} MB, "
                    f"over the {MAX_UNPACKED_BYTES // (1024 * 1024)} MB limit")
        return None

    @staticmethod
    def _fail(result, document_status, note, category, started):
        result.status = "error" if document_status != DocumentStatus.UNSUPPORTED else "skipped"
        result.document_status = document_status
        result.note = note
        result.errors.append(ConversionError(message=note[:240], category=category))
        result.timings_ms = {"total": (time.monotonic() - started) * 1000.0}
        return result

    # -- formats -------------------------------------------------------
    def _pdf(self, path: Path, limit: int):
        import pypdf

        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:  # noqa: BLE001
                unlocked = 0
            if not unlocked:
                raise _Refused(DocumentStatus.CORRUPT, "the PDF is password-protected", "policy")
        total_pages = len(reader.pages)
        parts: list[str] = []
        size = 0
        truncated = False
        for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            page_text = (page.extract_text() or "").strip()
            if not page_text:
                continue
            block = f"[page {number}]\n{page_text}"
            if size + len(block) > limit:
                block = block[: max(0, limit - size)]
                truncated = True
            parts.append(block)
            size += len(block)
            if truncated:
                break
        text = "\n\n".join(parts)
        if not text.strip():
            # A scan: the pages exist but carry no text layer. This tier has no OCR.
            return "", total_pages, "the PDF has no text layer (a scan); OCR is not available here", (
                "skipped", DocumentStatus.OCR_REQUIRED)
        notes = []
        if truncated:
            notes.append(f"text truncated at {limit:,} chars")
        if total_pages > MAX_PDF_PAGES:
            notes.append(f"only the first {MAX_PDF_PAGES} of {total_pages} pages were read")
        return text, total_pages, "; ".join(notes), (("partial", DocumentStatus.PARTIAL) if notes else None)

    def _docx(self, path: Path, limit: int):
        import docx

        document = docx.Document(str(path))
        parts: list[str] = []
        headings: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            parts.append(text)
            if (paragraph.style.name or "").lower().startswith("heading"):
                headings.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts)
        note, status = "", None
        if len(text) > limit:
            text, note, status = text[:limit], f"text truncated at {limit:,} chars", ("partial", DocumentStatus.PARTIAL)
        return text, len(document.paragraphs), note, status

    def _xlsx(self, path: Path, limit: int):
        import openpyxl

        # read_only streams rows; data_only returns cached values, not formulas.
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        size = 0
        truncated_rows = False
        truncated_chars = False
        try:
            sheets = workbook.sheetnames
            for name in sheets:
                sheet = workbook[name]
                rows: list[str] = []
                for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if index > MAX_XLSX_ROWS_PER_SHEET:
                        truncated_rows = True
                        break
                    cells = ["" if value is None else str(value).strip() for value in row]
                    while cells and not cells[-1]:
                        cells.pop()
                    if any(cells):
                        rows.append(" | ".join(cells))
                block = f"## Sheet: {name}\n" + "\n".join(rows)
                if size + len(block) > limit:
                    block = block[: max(0, limit - size)]
                    truncated_chars = True
                parts.append(block)
                size += len(block)
                if truncated_chars:
                    break
        finally:
            workbook.close()
        text = "\n\n".join(parts)
        # A workbook of empty sheets still produces headings; count real content.
        has_content = any("\n" in p and p.split("\n", 1)[1].strip() for p in parts)
        if not has_content:
            return "", len(sheets), "the workbook has no cell values (formulas without cached values read as empty)", (
                "empty", DocumentStatus.PARTIAL)
        notes = []
        if truncated_chars:
            notes.append(f"text truncated at {limit:,} chars")
        if truncated_rows:
            notes.append(f"only the first {MAX_XLSX_ROWS_PER_SHEET:,} rows of a sheet were read")
        return text, len(sheets), "; ".join(notes), (("partial", DocumentStatus.PARTIAL) if notes else None)


class _Refused(Exception):
    """A typed refusal raised inside an extractor."""

    def __init__(self, status: DocumentStatus, message: str, category: str) -> None:
        super().__init__(message)
        self.status = status
        self.category = category
